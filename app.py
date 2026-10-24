
import json
import os
from datetime import datetime

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

load_dotenv()

st.set_page_config(
    page_title="AI Compliance Workbench",
    page_icon="🛡️",
    layout="wide",
)

st.title("🛡️ AI Compliance Workbench")
st.caption(
    "An AI-assisted workspace for compliance evidence review. "
    "AI findings are suggestions, not legal advice."
)

SUPPORTED_TYPES = ["pdf", "docx", "xlsx", "jpg", "jpeg"]
MAX_FILE_SIZE_MB = 20
MAX_CONTEXT_CHARS = 45000


# -------------------------
# DOCUMENT TEXT EXTRACTION
# -------------------------

def extract_pdf(uploaded_file):
    reader = PdfReader(uploaded_file)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""

        if page_text.strip():
            pages.append(f"[Page {page_number}]\n{page_text}")

    return "\n\n".join(pages)


def extract_docx(uploaded_file):
    document = Document(uploaded_file)

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table_index, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            rows.append(
                " | ".join(cell.text.strip() for cell in row.cells)
            )

        parts.append(
            f"[Table {table_index}]\n" + "\n".join(rows)
        )

    return "\n\n".join(parts)


def extract_xlsx(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):
            values = [
                "" if value is None else str(value)
                for value in row
            ]

            if any(value.strip() for value in values):
                parts.append(" | ".join(values))

    workbook.close()
    return "\n".join(parts)


def extract_text(uploaded_file):
    filename = uploaded_file.name.lower()
    uploaded_file.seek(0)

    if filename.endswith(".pdf"):
        return extract_pdf(uploaded_file)

    if filename.endswith(".docx"):
        return extract_docx(uploaded_file)

    if filename.endswith(".xlsx"):
        return extract_xlsx(uploaded_file)

    # Image OCR is not implemented in this version.
    return ""


# -------------------------
# AI CONFIGURATION
# -------------------------

def get_openai_client():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()

    if not api_key:
        return None

    try:
        from openai import OpenAI
        return OpenAI(api_key=api_key)
    except Exception:
        return None


def build_context(documents, selected_names=None):
    chunks = []

    for item in documents:
        if (
            selected_names is not None
            and item["name"] not in selected_names
        ):
            continue

        body = item.get("text", "")

        if body.strip():
            chunks.append(
                f"DOCUMENT: {item['name']}\n{body}"
            )

    return "\n\n---\n\n".join(chunks)[:MAX_CONTEXT_CHARS]


def ask_ai(prompt, context, system_instruction):
    client = get_openai_client()

    if client is None:
        return None, (
            "AI is not configured. Add OPENAI_API_KEY to your .env "
            "file and restart the application."
        )

    if not context.strip():
        return None, (
            "No extractable text was found in the selected documents. "
            "Scanned PDFs and images may require OCR."
        )

    try:
        response = client.chat.completions.create(
            model=os.getenv(
                "OPENAI_MODEL",
                "gpt-4.1-mini",
            ),
            temperature=0.1,
            messages=[
                {
                    "role": "system",
                    "content": system_instruction,
                },
                {
                    "role": "user",
                    "content": (
                        f"DOCUMENT EVIDENCE:\n{context}\n\n"
                        f"REQUEST:\n{prompt}"
                    ),
                },
            ],
        )

        return response.choices[0].message.content, None

    except Exception as exc:
        return None, f"AI request failed: {exc}"


# -------------------------
# SESSION STORAGE
# -------------------------

if "documents" not in st.session_state:
    st.session_state.documents = {}


# -------------------------
# SIDEBAR
# -------------------------

with st.sidebar:
    st.header("Workspace")

    st.write(
        "Upload policies, procedures, and compliance evidence."
    )

    st.caption("Accepted: PDF, DOCX, XLSX, JPG, JPEG")
    st.caption(
        f"Maximum file size: {MAX_FILE_SIZE_MB} MB per file"
    )

    st.divider()

    if os.getenv("OPENAI_API_KEY", "").strip():
        st.success("AI is configured")
    else:
        st.info(
            "AI is not configured. Document upload and text "
            "extraction can still work."
        )


# -------------------------
# FILE UPLOAD
# -------------------------

upload_col, privacy_col = st.columns([2, 1])

with upload_col:
    uploaded_files = st.file_uploader(
        "Upload compliance documents",
        type=SUPPORTED_TYPES,
        accept_multiple_files=True,
        help="Files are processed during this app session.",
    )

with privacy_col:
    st.markdown("**Privacy note**")

    st.write(
        "Files remain in the app session unless extracted text "
        "is sent to an AI provider. Use non-sensitive demo files."
    )


if uploaded_files:
    for uploaded in uploaded_files:

        if uploaded.size > MAX_FILE_SIZE_MB * 1024 * 1024:
            st.warning(
                f"{uploaded.name} exceeds the "
                f"{MAX_FILE_SIZE_MB} MB limit and was skipped."
            )
            continue

        if uploaded.name not in st.session_state.documents:
            try:
                extracted = extract_text(uploaded)

                if extracted.strip():
                    status = "Text extracted"
                else:
                    status = (
                        "No text extracted: image, scanned file, "
                        "or empty document"
                    )

            except Exception as exc:
                extracted = ""
                status = f"Extraction failed: {exc}"

            st.session_state.documents[uploaded.name] = {
                "name": uploaded.name,
                "type": uploaded.name.rsplit(".", 1)[-1].lower(),
                "size": uploaded.size,
                "text": extracted,
                "uploaded_at": datetime.now().strftime(
                    "%Y-%m-%d %H:%M"
                ),
                "extraction_status": status,
            }


documents = list(st.session_state.documents.values())


# -------------------------
# DASHBOARD METRICS
# -------------------------

if documents:
    total_size_mb = (
        sum(item["size"] for item in documents)
        / (1024 * 1024)
    )

    readable_count = sum(
        bool(item["text"].strip())
        for item in documents
    )

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric(
        "Documents in session",
        len(documents),
    )

    metric2.metric(
        "Documents with extracted text",
        readable_count,
    )

    metric3.metric(
        "Total upload size",
        f"{total_size_mb:.1f} MB",
    )


tab_docs, tab_ask, tab_review, tab_export = st.tabs(
    [
        " Document Library",
        " Ask Documents",
        " Compliance Review",
        " Export Report",
    ]
)




with tab_docs:
    st.subheader("Document Library")

    if not documents:
        st.info("Upload one or more files above to get started.")

    else:
        search = st.text_input(
            "Search by filename",
            placeholder="Example: access policy, incident, report",
        )

        filtered = [
            item
            for item in documents
            if search.strip().lower() in item["name"].lower()
        ]

        st.caption(
            f"Showing {len(filtered)} of {len(documents)} documents"
        )

        if filtered:
            rows = [
                {
                    "Filename": item["name"],
                    "Type": item["type"].upper(),
                    "Size (KB)": round(
                        item["size"] / 1024, 1
                    ),
                    "Text extraction": item["extraction_status"],
                }
                for item in filtered
            ]

            st.dataframe(
                pd.DataFrame(rows),
                use_container_width=True,
                hide_index=True,
            )

            selected_name = st.selectbox(
                "Preview a document",
                [item["name"] for item in filtered],
            )

            selected_doc = next(
                item
                for item in filtered
                if item["name"] == selected_name
            )

            st.caption(selected_doc["extraction_status"])

            if selected_doc["text"].strip():
                with st.expander("Preview extracted text"):
                    st.text(
                        selected_doc["text"][:12000]
                    )

            elif selected_doc["type"] in ("jpg", "jpeg"):
                st.info(
                    "Image OCR is not enabled in this version."
                )

            else:
                st.info(
                    "No readable text was extracted. "
                    "The file may be scanned or empty."
                )

        if st.button("Clear all uploaded documents"):
            st.session_state.documents = {}
            st.rerun()


# -------------------------
# TAB 2: ASK DOCUMENTS
# -------------------------

with tab_ask:
    st.subheader("Ask Questions About Your Documents")

    st.write(
        "Ask questions and receive answers grounded in the "
        "extracted document text."
    )

    if not documents:
        st.info("Upload documents first.")

    else:
        available_docs = [
            item["name"]
            for item in documents
            if item["text"].strip()
        ]

        selected_for_qa = st.multiselect(
            "Select documents to use as evidence",
            options=available_docs,
            default=available_docs,
        )

        question = st.text_area(
            "Your question",
            placeholder=(
                "Example: What does the policy say "
                "about incident reporting?"
            ),
            height=90,
        )

        if st.button(
            "Answer from documents",
            type="primary",
            key="ask_button",
        ):
            context = build_context(
                documents,
                selected_for_qa,
            )

            if not question.strip():
                st.warning("Enter a question first.")

            else:
                with st.spinner(
                    "Reviewing the selected document evidence..."
                ):
                    answer, error = ask_ai(
                        question,
                        context,
                        (
                            "You are a careful compliance document "
                            "assistant. Answer only from supplied "
                            "evidence. If the answer is absent, say so. "
                            "Cite filenames and page labels when available. "
                            "Separate evidence from interpretation. "
                            "Do not claim legal compliance or provide "
                            "legal advice."
                        ),
                    )

                if error:
                    st.error(error)

                else:
                    st.markdown(answer)

                    st.session_state["last_answer"] = answer
                    st.session_state["last_question"] = question


# -------------------------
# TAB 3: COMPLIANCE REVIEW
# -------------------------

with tab_review:
    st.subheader("AI-Assisted Compliance Review")

    st.write(
        "This is a preliminary evidence review, not a legal "
        "determination. A human reviewer must confirm all findings."
    )

    if not documents:
        st.info("Upload documents first.")

    else:
        checklist = [
            "Access control and user permissions",
            "Information security responsibilities",
            "Security policy review and approval",
            "Incident reporting and response",
            "Data protection and retention",
            "Supplier or third-party risk management",
            "Training and awareness",
            "Monitoring, logging, and audit evidence",
        ]

        review_items = st.multiselect(
            "Select review areas",
            checklist,
            default=checklist[:5],
        )

        available_review_docs = [
            item["name"]
            for item in documents
            if item["text"].strip()
        ]

        review_docs = st.multiselect(
            "Documents to review",
            options=available_review_docs,
            default=available_review_docs,
            key="review_docs",
        )

        if st.button(
            "Run preliminary review",
            type="primary",
        ):
            context = build_context(
                documents,
                review_docs,
            )

            if not review_items:
                st.warning("Select at least one review area.")

            else:
                request = (
                    "Review the supplied document evidence against "
                    "these areas:\n"
                    + "\n".join(
                        f"- {item}" for item in review_items
                    )
                    + "\n\nFor each area, provide:\n"
                    "- Area\n"
                    "- Status: Evidence found, Potential gap, "
                    "or Not enough information\n"
                    "- Supporting evidence with filename and page "
                    "when available\n"
                    "- Rationale\n"
                    "- Recommended follow-up\n"
                    "- Confidence: Low, Medium, or High\n\n"
                    "Missing evidence does not prove non-compliance. "
                    "Do not invent quotations, page numbers, or obligations."
                )

                with st.spinner("Reviewing compliance evidence..."):
                    review, error = ask_ai(
                        request,
                        context,
                        (
                            "You are an evidence-focused compliance "
                            "analyst. Be cautious, transparent, and "
                            "specific. Do not provide legal advice. "
                            "Never invent evidence or claim compliance "
                            "without supporting evidence."
                        ),
                    )

                if error:
                    st.error(error)

                else:
                    st.markdown(review)

                    st.session_state["last_review"] = review
                    st.session_state["last_review_areas"] = review_items


# -------------------------
# TAB 4: EXPORT REPORT
# -------------------------

with tab_export:
    st.subheader("Export Review Notes")

    st.write(
        "Download a report containing document filenames "
        "and generated review notes."
    )

    report = {
        "project": "AI Compliance Workbench",
        "generated_at": datetime.now().isoformat(
            timespec="minutes"
        ),
        "documents": [
            {
                "filename": item["name"],
                "type": item["type"],
                "size_bytes": item["size"],
                "text_extraction": item["extraction_status"],
            }
            for item in documents
        ],
        "last_question": st.session_state.get(
            "last_question", ""
        ),
        "last_answer": st.session_state.get(
            "last_answer", ""
        ),
        "last_review": st.session_state.get(
            "last_review", ""
        ),
        "review_areas": st.session_state.get(
            "last_review_areas", []
        ),
        "disclaimer": (
            "AI-generated notes are preliminary, require "
            "human verification, and are not legal advice."
        ),
    }

    st.download_button(
        "Download session report (JSON)",
        data=json.dumps(
            report,
            indent=2,
            ensure_ascii=False,
        ),
        file_name="compliance_workbench_report.json",
        mime="application/json",
    )

    markdown_report = (
        "# AI Compliance Workbench — Session Report\n\n"
        f"Generated: {report['generated_at']}\n\n"
        "## Documents\n"
        + "\n".join(
            f"- **{item['filename']}** — "
            f"{item['type'].upper()} — "
            f"{item['text_extraction']}"
            for item in report["documents"]
        )
        + "\n\n## Last question\n"
        + (
            report["last_question"]
            or "_No question submitted._"
        )
        + "\n\n## Last answer\n"
        + (
            report["last_answer"]
            or "_No answer generated._"
        )
        + "\n\n## Last compliance review\n"
        + (
            report["last_review"]
            or "_No review generated._"
        )
        + (
            "\n\n> AI-generated content is preliminary, "
            "must be verified by a human, and is not legal advice.\n"
        )
    )

    st.download_button(
        "Download session report (Markdown)",
        data=markdown_report,
        file_name="compliance_workbench_report.md",
        mime="text/markdown",
    )


# -------------------------
# FOOTER
# -------------------------

st.divider()

st.caption(
    "Prototype limitations: session-based storage, no database, "
    "no authentication, no audit logging, and no image OCR. "
    "Do not upload confidential company data to this demo."
)
