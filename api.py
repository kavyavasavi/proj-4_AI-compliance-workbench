import hashlib
import io
import json
import os
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path
from threading import Lock

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from google import genai
from google.genai import types

load_dotenv()
app = FastAPI(title="AI Compliance Workbench API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "https://proj-4-ai-compliance-workbench.vercel.app",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["*"],
)

SUPPORTED_TYPES = {"pdf", "docx", "xlsx"}
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_CONTEXT_CHARS = 45000
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "compliance_workbench.db"

# In-memory cache of extracted document text; SQLite is persistent storage.
sessions = {}
sessions_lock = Lock()


def get_db_connection():
    connection = sqlite3.connect(DB_PATH, timeout=30)
    connection.row_factory = sqlite3.Row
    return connection


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def make_short_title(text, max_length=48):
    """Create a readable sidebar title from a filename or user question."""
    value = (text or "").strip()
    # Remove a file extension when the source is a filename.
    if "." in value and value.rsplit(".", 1)[-1].lower() in {
        "pdf", "docx", "xlsx", "txt"
    }:
        value = value.rsplit(".", 1)[0]
    value = value.replace("_", " ").replace("-", " ")
    value = " ".join(value.split())
    if not value:
        return "New Compliance Review"
    if len(value) > max_length:
        value = value[: max_length - 1].rsplit(" ", 1)[0] or value[: max_length - 1]
        value += "…"
    return value[:1].upper() + value[1:]


def update_title_if_default(session_id, source_text):
    """Rename only generic sessions; never overwrite a title chosen by the user."""
    title = make_short_title(source_text)
    if not title or title == "New Compliance Review":
        return
    with get_db_connection() as connection:
        row = connection.execute(
            "SELECT title FROM sessions WHERE id = ?", (session_id,)
        ).fetchone()
        if row and row["title"].strip().lower() in {
            "new compliance review", "new session", "untitled session"
        }:
            connection.execute(
                "UPDATE sessions SET title = ?, updated_at = ? WHERE id = ?",
                (title, now_iso(), session_id),
            )


def initialize_database():
    with get_db_connection() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS sessions (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL DEFAULT 'New Compliance Review',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'active'
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS documents (
                session_id TEXT NOT NULL,
                filename TEXT NOT NULL,
                document_json TEXT NOT NULL,
                PRIMARY KEY (session_id, filename)
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id TEXT NOT NULL,
                kind TEXT NOT NULL,
                created_at TEXT NOT NULL,
                payload_json TEXT NOT NULL
            )
        """)


def load_saved_documents():
    with get_db_connection() as connection:
        rows = connection.execute(
            "SELECT session_id, filename, document_json FROM documents"
        ).fetchall()
    with sessions_lock:
        for row in rows:
            sessions.setdefault(row["session_id"], {})[row["filename"]] = json.loads(
                row["document_json"]
            )


def get_documents(session_id):
    with sessions_lock:
        return sessions.setdefault(session_id, {})


def ensure_session_exists(session_id):
    with get_db_connection() as connection:
        row = connection.execute(
            "SELECT id FROM sessions WHERE id = ?", (session_id,)
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Session not found.")


def save_session_documents(session_id):
    with sessions_lock:
        documents = dict(sessions.get(session_id, {}))
    with get_db_connection() as connection:
        connection.execute("DELETE FROM documents WHERE session_id = ?", (session_id,))
        connection.executemany(
            """
            INSERT INTO documents (session_id, filename, document_json)
            VALUES (?, ?, ?)
            """,
            [
                (session_id, filename, json.dumps(document, ensure_ascii=False))
                for filename, document in documents.items()
            ],
        )
        connection.execute(
            "UPDATE sessions SET updated_at = ? WHERE id = ?",
            (now_iso(), session_id),
        )


def save_history(session_id, kind, payload):
    created_at = now_iso()
    with get_db_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO history (session_id, kind, created_at, payload_json)
            VALUES (?, ?, ?, ?)
            """,
            (session_id, kind, created_at, json.dumps(payload, ensure_ascii=False)),
        )
        connection.execute(
            "UPDATE sessions SET updated_at = ? WHERE id = ?",
            (created_at, session_id),
        )
        return {"id": cursor.lastrowid, "created_at": created_at}


def get_session_history(session_id):
    ensure_session_exists(session_id)
    with get_db_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, kind, created_at, payload_json
            FROM history WHERE session_id = ? ORDER BY id ASC
            """,
            (session_id,),
        ).fetchall()
    result = []
    for row in rows:
        try:
            payload = json.loads(row["payload_json"])
        except json.JSONDecodeError:
            payload = {}
        result.append({
            "id": row["id"],
            "kind": row["kind"],
            "created_at": row["created_at"],
            "payload": payload,
        })
    return result


def list_documents(session_id):
    documents = get_documents(session_id)
    return {
        "documents": [
            {
                "name": item["name"],
                "type": item["type"],
                "size_bytes": item["size"],
                "uploaded_at": item["uploaded_at"],
                "extraction_status": item["extraction_status"],
                "has_text": bool(item.get("text", "").strip()),
            }
            for item in documents.values()
        ]
    }


def extract_document(filename, file_bytes):
    stream = io.BytesIO(file_bytes)
    name = filename.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(stream)
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {number}]\n{text}")
        return "\n\n".join(pages)

    if name.endswith(".docx"):
        document = Document(stream)
        parts = [
            paragraph.text for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for index, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            parts.append(f"[Table {index}]\n" + "\n".join(rows))
        return "\n\n".join(parts)

    if name.endswith(".xlsx"):
        workbook = load_workbook(stream, read_only=True, data_only=True)
        parts = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        parts.append(" | ".join(values))
        finally:
            workbook.close()
        return "\n".join(parts)

    raise ValueError("Unsupported file type.")


def build_context(documents, selected_names=None):
    chunks = []
    for name, item in documents.items():
        if selected_names is not None and name not in selected_names:
            continue
        if item.get("text", "").strip():
            chunks.append(f"DOCUMENT: {name}\n{item['text']}")
    return "\n\n---\n\n".join(chunks)[:MAX_CONTEXT_CHARS]


def ask_ai(prompt, context, system_instruction, conversation_history=None):
    """Call Gemini with optional evidence and a short, session-local chat history."""
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise HTTPException(
            status_code=503, detail="Gemini is not configured on the server."
        )

    evidence = (context or "").strip()
    history_text = ""
    if conversation_history:
        history_lines = []
        for item in conversation_history[-8:]:
            payload = item.get("payload") or {}
            question = str(payload.get("question") or "").strip()
            answer = str(payload.get("answer") or "").strip()
            if question and answer:
                history_lines.append(f"User: {question}\nAssistant: {answer}")
        if history_lines:
            history_text = "RECENT CONVERSATION (for follow-up context only):\n" + "\n\n".join(history_lines)

    if evidence:
        contents = f"DOCUMENT EVIDENCE:\n{evidence}\n\n"
    else:
        contents = (
            "DOCUMENT EVIDENCE: None was selected or available. Do not pretend to have read documents. "
            "For document-specific or compliance-specific questions, explain that the user must upload/select readable evidence. "
            "You may answer greetings and general conversational questions briefly.\n\n"
        )
    if history_text:
        contents += history_text + "\n\n"
    contents += f"CURRENT USER MESSAGE:\n{prompt}"

    try:
        client = genai.Client(api_key=api_key)
        model = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite").strip()
        response = client.models.generate_content(
            model=model,
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=system_instruction,
                temperature=0.2,
            ),
        )
        response_text = getattr(response, "text", None)
        if not response_text or not response_text.strip():
            raise HTTPException(status_code=502, detail="Gemini returned an empty response.")
        return response_text.strip()
    except HTTPException:
        raise
    except Exception as exc:
        # Keep the API response useful without exposing secrets or provider internals.
        raise HTTPException(
            status_code=502,
            detail="The AI request failed. Check the backend terminal and Gemini model/API configuration.",
        ) from exc


class CreateSessionRequest(BaseModel):
    title: str | None = None


class RenameSessionRequest(BaseModel):
    title: str


class QuestionRequest(BaseModel):
    session_id: str
    question: str
    document_names: list[str] | None = None


class ReviewRequest(BaseModel):
    session_id: str
    review_areas: list[str]
    document_names: list[str] | None = None


initialize_database()
load_saved_documents()


@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "AI Compliance Workbench API",
        "gemini_configured": bool(os.getenv("GEMINI_API_KEY", "").strip()),
    }


@app.post("/sessions")
def create_session(request: CreateSessionRequest):
    session_id = str(uuid.uuid4())
    timestamp = now_iso()
    title = (request.title or "").strip() or "New Compliance Review"
    title = title[:120]
    with get_db_connection() as connection:
        connection.execute(
            """
            INSERT INTO sessions (id, title, created_at, updated_at, status)
            VALUES (?, ?, ?, ?, 'active')
            """,
            (session_id, title, timestamp, timestamp),
        )
    with sessions_lock:
        sessions.setdefault(session_id, {})
    return {
        "id": session_id,
        "title": title,
        "created_at": timestamp,
        "updated_at": timestamp,
        "status": "active",
    }


@app.get("/sessions")
def list_sessions():
    with get_db_connection() as connection:
        rows = connection.execute(
            """
            SELECT s.id, s.title, s.created_at, s.updated_at, s.status,
                   COUNT(DISTINCT d.filename) AS document_count,
                   COUNT(DISTINCT h.id) AS history_count
            FROM sessions s
            LEFT JOIN documents d ON d.session_id = s.id
            LEFT JOIN history h ON h.session_id = s.id
            GROUP BY s.id
            ORDER BY s.updated_at DESC, s.created_at DESC
            """
        ).fetchall()
    return {"sessions": [dict(row) for row in rows]}


@app.get("/sessions/{session_id}")
def get_session(session_id: str):
    ensure_session_exists(session_id)
    with get_db_connection() as connection:
        row = connection.execute(
            """
            SELECT id, title, created_at, updated_at, status
            FROM sessions WHERE id = ?
            """,
            (session_id,),
        ).fetchone()
    return {
        **dict(row),
        **list_documents(session_id),
        "history": get_session_history(session_id),
    }


@app.patch("/sessions/{session_id}")
def rename_session(session_id: str, request: RenameSessionRequest):
    title = request.title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Session title cannot be empty.")
    title = title[:120]
    ensure_session_exists(session_id)
    updated_at = now_iso()
    with get_db_connection() as connection:
        connection.execute(
            "UPDATE sessions SET title = ?, updated_at = ? WHERE id = ?",
            (title, updated_at, session_id),
        )
    return {"id": session_id, "title": title, "updated_at": updated_at}


@app.delete("/sessions/{session_id}")
def delete_session(session_id: str):
    ensure_session_exists(session_id)
    with sessions_lock:
        sessions.pop(session_id, None)
    with get_db_connection() as connection:
        connection.execute("DELETE FROM documents WHERE session_id = ?", (session_id,))
        connection.execute("DELETE FROM history WHERE session_id = ?", (session_id,))
        connection.execute("DELETE FROM sessions WHERE id = ?", (session_id,))
    return {"message": "Session and its saved data deleted."}


@app.post("/documents")
async def upload_documents(
    session_id: str = Form(...),
    files: list[UploadFile] = File(...),
):
    if not session_id.strip():
        raise HTTPException(status_code=400, detail="Session ID required.")
    ensure_session_exists(session_id)
    if not files:
        raise HTTPException(status_code=400, detail="Choose files to upload.")

    documents = get_documents(session_id)
    results = []
    for uploaded in files:
        filename = os.path.basename(uploaded.filename or "").strip()
        if not filename or "." not in filename:
            results.append({
                "filename": filename or "unknown",
                "status": "failed",
                "message": "Invalid filename.",
            })
            continue

        extension = filename.rsplit(".", 1)[-1].lower()
        if extension not in SUPPORTED_TYPES:
            results.append({
                "filename": filename,
                "status": "failed",
                "message": "Supported formats: PDF, DOCX, XLSX.",
            })
            continue

        file_bytes = await uploaded.read()
        if len(file_bytes) > MAX_FILE_SIZE:
            results.append({
                "filename": filename,
                "status": "failed",
                "message": "File exceeds the 20 MB limit.",
            })
            continue

        file_hash = hashlib.sha256(file_bytes).hexdigest()
        if filename in documents:
            existing_hash = documents[filename].get("sha256")
            message = (
                "This document has already been uploaded."
                if existing_hash == file_hash
                else (
                    "A document with this filename already exists. "
                    "Rename the new file if it is a different document."
                )
            )
            results.append({
                "filename": filename,
                "status": "duplicate",
                "message": message,
            })
            continue

        duplicate = next(
            (doc for doc in documents.values() if doc.get("sha256") == file_hash),
            None,
        )
        if duplicate:
            results.append({
                "filename": filename,
                "status": "duplicate",
                "message": (
                    f"This file has the same content as '{duplicate['name']}', "
                    "which is already uploaded."
                ),
            })
            continue

        try:
            extracted = extract_document(filename, file_bytes)
            status = "Text extracted" if extracted.strip() else "No extractable text found"
            documents[filename] = {
                "name": filename,
                "type": extension,
                "size": len(file_bytes),
                "text": extracted,
                "sha256": file_hash,
                "uploaded_at": datetime.now().isoformat(timespec="minutes"),
                "extraction_status": status,
            }
            results.append({
                "filename": filename,
                "status": "success",
                "extraction_status": status,
                "size_bytes": len(file_bytes),
            })
        except Exception:
            results.append({
                "filename": filename,
                "status": "failed",
                "message": "Could not read this file. Check its format and contents.",
            })

    save_session_documents(session_id)

    # Give a generic session a useful title based on the first successfully
    # uploaded document. Existing/custom titles are preserved.
    first_uploaded = next(
        (item["filename"] for item in results if item.get("status") == "success"),
        None,
    )
    if first_uploaded:
        update_title_if_default(session_id, first_uploaded)

    return {
        "documents": results,
        "library": list_documents(session_id)["documents"],
    }


@app.get("/documents/{session_id}")
def get_document_library(session_id: str):
    ensure_session_exists(session_id)
    return list_documents(session_id)


@app.delete("/documents/{session_id}")
def clear_documents(session_id: str):
    ensure_session_exists(session_id)
    with sessions_lock:
        sessions[session_id] = {}
    with get_db_connection() as connection:
        connection.execute("DELETE FROM documents WHERE session_id = ?", (session_id,))
        connection.execute(
            "UPDATE sessions SET updated_at = ? WHERE id = ?",
            (now_iso(), session_id),
        )
    return {"message": "Session documents cleared."}


@app.post("/ask")
def ask_documents(request: QuestionRequest):
    if not request.question.strip():
        raise HTTPException(status_code=400, detail="Enter a question.")
    ensure_session_exists(request.session_id)
    documents = get_documents(request.session_id)
    if request.document_names is not None:
        missing = set(request.document_names) - set(documents)
        if missing:
            raise HTTPException(
                status_code=404,
                detail="One or more selected documents were not found.",
            )

    context = build_context(documents, request.document_names)
    previous_qa = [
        item for item in get_session_history(request.session_id)
        if item.get("kind") == "qa"
    ][-8:]
    answer = ask_ai(
        request.question.strip(),
        context,
        (
            "You are the AI assistant inside a compliance document workbench. "
            "Respond naturally and briefly to greetings and ordinary conversational messages. "
            "For document-specific questions, use only the supplied document evidence and relevant prior conversation. "
            "If evidence is absent or insufficient, say so clearly; never invent document content, quotations, filenames, page numbers, or regulatory obligations. "
            "When possible, identify the filename supporting a claim. Separate evidence from interpretation. "
            "Do not claim that an organization is compliant and do not provide legal advice. "
            "Prior conversation is context, not proof of what a document says."
        ),
        conversation_history=previous_qa,
    )
    payload = {
        "question": request.question,
        "answer": answer,
        "document_names": request.document_names,
    }
    record = save_history(request.session_id, "qa", payload)
    update_title_if_default(request.session_id, request.question)
    return {
        "id": record["id"],
        "created_at": record["created_at"],
        **payload,
    }


@app.post("/review")
def compliance_review(request: ReviewRequest):
    if not request.review_areas:
        raise HTTPException(status_code=400, detail="Select at least one review area.")
    ensure_session_exists(request.session_id)
    documents = get_documents(request.session_id)
    if request.document_names is not None:
        missing = set(request.document_names) - set(documents)
        if missing:
            raise HTTPException(
                status_code=404,
                detail="One or more selected documents were not found.",
            )

    context = build_context(documents, request.document_names)
    prompt = (
        "Review the supplied document evidence against these areas:\n"
        + "\n".join(f"- {area}" for area in request.review_areas)
        + "\n\nFor each area, provide:\n"
        "- Area\n"
        "- Status: Evidence found, Potential gap, or Not enough information\n"
        "- Supporting evidence with filename and page when available\n"
        "- Rationale\n"
        "- Recommended follow-up\n"
        "- Confidence: Low, Medium, or High\n\n"
        "Missing evidence does not prove non-compliance. "
        "Do not invent quotations, page numbers, or obligations."
    )
    review = ask_ai(
        prompt,
        context,
        (
            "You are an evidence-focused compliance analyst. "
            "Be cautious, transparent, and specific. Do not provide "
            "legal advice. Never invent evidence or claim compliance "
            "without supporting evidence."
        ),
    )
    payload = {
        "review": review,
        "review_areas": request.review_areas,
        "document_names": request.document_names,
        "disclaimer": (
            "Preliminary AI findings require human verification "
            "and are not legal advice."
        ),
    }
    record = save_history(request.session_id, "review", payload)
    return {
        "id": record["id"],
        "created_at": record["created_at"],
        **payload,
    }


@app.get("/report/{session_id}")
def export_report(session_id: str):
    ensure_session_exists(session_id)
    documents = get_documents(session_id)
    report = {
        "project": "AI Compliance Workbench",
        "generated_at": datetime.now().isoformat(timespec="minutes"),
        "session_id": session_id,
        "documents": [
            {
                "filename": item["name"],
                "type": item["type"],
                "size_bytes": item["size"],
                "text_extraction": item["extraction_status"],
            }
            for item in documents.values()
        ],
        "history": get_session_history(session_id),
        "disclaimer": (
            "AI-generated notes are preliminary, require human "
            "verification, and are not legal advice."
        ),
    }
    return JSONResponse(content=report)
